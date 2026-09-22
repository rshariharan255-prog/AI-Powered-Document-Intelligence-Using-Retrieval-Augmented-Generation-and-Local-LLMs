import io
from pptx import Presentation
from pptx.util import Inches, Pt
from fastapi.testclient import TestClient
from main import app

def create_sample_ml_presentation() -> bytes:
    """Create a sample 3-slide PowerPoint presentation in memory."""
    prs = Presentation()
    
    # Slide 1: Deep Learning Architectures
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "Deep Learning Overview"
    tf1 = slide1.shapes.placeholders[1].text_frame
    tf1.text = "Convolutional Neural Networks (CNNs):"
    p1 = tf1.add_paragraph()
    p1.text = "- Used for spatial and image feature extraction"
    p2 = tf1.add_paragraph()
    p2.text = "- Key layers: Convolution, ReLU activation, Max Pooling, and Fully Connected layers"

    # Slide 2: Recurrent Networks
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "Sequence Modeling with RNNs"
    tf2 = slide2.shapes.placeholders[1].text_frame
    tf2.text = "Recurrent Neural Networks (RNNs) and LSTMs:"
    p3 = tf2.add_paragraph()
    p3.text = "- Designed for sequential and time-series data"
    p4 = tf2.add_paragraph()
    p4.text = "- LSTMs mitigate the vanishing gradient problem using Forget, Input, and Output gates"

    # Slide 3: Transformers
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    slide3.shapes.title.text = "Transformer Architecture"
    tf3 = slide3.shapes.placeholders[1].text_frame
    tf3.text = "Attention Is All You Need:"
    p5 = tf3.add_paragraph()
    p5.text = "- Self-Attention mechanism computes relationships between all tokens in parallel"
    p6 = tf3.add_paragraph()
    p6.text = "- Forms the foundation for modern Large Language Models like Gemma, BERT, and GPT"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

def run_presentation_test():
    print("=" * 65)
    print("POWERPOINT & PRESENTATION RAG PIPELINE TEST")
    print("=" * 65)

    client = TestClient(app)

    # 1. Upload PPTX
    print("\n[1/3] Uploading PowerPoint Presentation (POST /upload)...")
    pptx_bytes = create_sample_ml_presentation()
    upload_res = client.post(
        "/upload",
        files={"file": ("Deep_Learning_Lecture.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    )
    assert upload_res.status_code == 200, f"Upload failed: {upload_res.text}"
    doc_info = upload_res.json()
    doc_id = doc_info["document_id"]
    print(f"[OK] Ingested: {doc_info['filename']} [ID: {doc_id}]")
    print(f"[OK] Total Slides: {doc_info['pages']}, Total Chunks: {doc_info['chunks']}")

    # 2. Test Question on Slide 1 (CNNs)
    print("\n[2/3] Asking Question about CNNs (POST /ask)...")
    q1 = "What are the main layers of a Convolutional Neural Network mentioned in the slides?"
    print(f"Question: \"{q1}\"")
    ask_res1 = client.post("/ask", json={"document_id": doc_id, "question": q1, "debug_mode": True})
    assert ask_res1.status_code == 200, f"Query failed: {ask_res1.text}"
    data1 = ask_res1.json()

    print("\n--- GEMMA 2B ANSWER ---")
    print(data1["answer"])
    print("-----------------------")
    print(f"[OK] Grounded: {data1['is_grounded']}")
    print(f"[OK] Cited Sources: {len(data1['sources'])} (Page/Slide: {[s['page'] for s in data1['sources']]})")
    assert data1["is_grounded"] is True
    assert any(s["page"] == 1 for s in data1["sources"])

    # 3. Test Question on Slide 2 (LSTMs)
    print("\n[3/3] Asking Question about LSTMs (POST /ask)...")
    q2 = "How do LSTMs solve the vanishing gradient problem according to the presentation?"
    print(f"Question: \"{q2}\"")
    ask_res2 = client.post("/ask", json={"document_id": doc_id, "question": q2})
    assert ask_res2.status_code == 200
    data2 = ask_res2.json()

    print("\n--- GEMMA 2B ANSWER ---")
    print(data2["answer"])
    print("-----------------------")
    assert data2["is_grounded"] is True
    assert any(s["page"] == 2 for s in data2["sources"])
    print("[OK] Successfully retrieved and answered from Slide 2!")

    print("\n" + "=" * 65)
    print("ALL PRESENTATION RAG TESTS PASSED FLAWLESSLY!")
    print("=" * 65)

if __name__ == "__main__":
    run_presentation_test()
