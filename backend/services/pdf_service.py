import logging
import re
import io
from pathlib import Path
from typing import List, Optional
from fastapi import HTTPException
from models.schemas import PageContent, PDFExtractionResult

try:
    import pymupdf
except ImportError:
    import fitz as pymupdf

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    Presentation = None
    MSO_SHAPE_TYPE = None

try:
    from rapidocr_onnxruntime import RapidOCR
    _ocr_engine = None
except ImportError:
    RapidOCR = None
    _ocr_engine = None

logger = logging.getLogger(__name__)

def get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None and RapidOCR is not None:
        try:
            _ocr_engine = RapidOCR()
            logger.info("RapidOCR engine initialized successfully.")
        except Exception as e:
            logger.error(f"Failed to initialize RapidOCR: {e}")
            _ocr_engine = None
    return _ocr_engine

def run_ocr_on_bytes(image_bytes: bytes) -> str:
    """Execute RapidOCR engine on raw image bytes."""
    if not image_bytes:
        return ""
    engine = get_ocr_engine()
    if engine is None:
        return ""
    try:
        results, _ = engine(image_bytes)
        if not results:
            return ""
        lines = [line[1].strip() for line in results if line and len(line) >= 2 and line[1].strip()]
        return "\n".join(lines)
    except Exception as e:
        logger.warning(f"OCR processing failed for image snippet: {e}")
        return ""

class DocumentParserService:
    @staticmethod
    def clean_text(text: str) -> str:
        """
        Perform gentle text cleaning:
        - Normalize multiple horizontal whitespace into a single space
        - Normalize multiple linebreaks
        - Preserve lists, bullet points, numbers, and technical terms
        """
        if not text:
            return ""
        
        text = text.replace('\xa0', ' ').replace('\t', ' ')
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
        text = re.sub(r'\n{3,}', '\n\n', text)
        lines = [re.sub(r' {2,}', ' ', line).strip() for line in text.splitlines()]
        cleaned = '\n'.join(lines).strip()
        return cleaned

    def _extract_image_file(self, file_path: Path) -> List[PageContent]:
        """Extract text from standalone image files (.png, .jpg, .jpeg, .webp, .bmp, .tiff) using OCR."""
        try:
            with open(file_path, "rb") as f:
                img_bytes = f.read()
            
            ocr_text = run_ocr_on_bytes(img_bytes)
            cleaned = self.clean_text(ocr_text)
            
            if not cleaned:
                cleaned = f"[Image: {file_path.name} (No readable text detected via OCR)]"

            return [
                PageContent(
                    page_number=1,
                    text=cleaned,
                    char_count=len(cleaned)
                )
            ]
        except Exception as e:
            logger.error(f"Failed to process image file {file_path}: {e}")
            raise HTTPException(
                status_code=400,
                detail=f"Unable to perform OCR on image file: {str(e)}"
            )

    def _extract_pdf(self, file_path: Path) -> List[PageContent]:
        """Extract text page-by-page from PDF with OCR fallback for scanned pages/images."""
        try:
            doc = pymupdf.open(str(file_path))
        except Exception as e:
            logger.error(f"Failed to open PDF file {file_path}: {str(e)}")
            raise HTTPException(
                status_code=400,
                detail="Unable to parse the PDF document. The file may be encrypted or corrupted."
            )

        pages: List[PageContent] = []
        try:
            for page_index in range(len(doc)):
                page = doc[page_index]
                raw_text = page.get_text("text") or ""
                cleaned_text = self.clean_text(raw_text)
                
                # If page has no readable digital text (scanned page/image PDF), run OCR
                needs_ocr = len(cleaned_text) < 30

                if needs_ocr:
                    try:
                        pix = page.get_pixmap(dpi=150)
                        img_bytes = pix.tobytes("png")
                        ocr_text = run_ocr_on_bytes(img_bytes)
                        cleaned_ocr = self.clean_text(ocr_text)
                        
                        if cleaned_ocr and cleaned_ocr not in cleaned_text:
                            if cleaned_text:
                                cleaned_text = f"{cleaned_text}\n\n[OCR Extracted Text from Page Image]:\n{cleaned_ocr}"
                            else:
                                cleaned_text = cleaned_ocr
                    except Exception as ocr_err:
                        logger.warning(f"Failed OCR on PDF page {page_index + 1}: {ocr_err}")

                page_number = page_index + 1
                char_count = len(cleaned_text)

                pages.append(
                    PageContent(
                        page_number=page_number,
                        text=cleaned_text,
                        char_count=char_count
                    )
                )
            return pages
        finally:
            doc.close()

    def _extract_pptx(self, file_path: Path) -> List[PageContent]:
        """Extract text slide-by-slide from PowerPoint (.pptx) with image OCR for diagrams/figures."""
        if Presentation is None:
            raise HTTPException(
                status_code=500,
                detail="python-pptx package is not installed on server."
            )

        try:
            prs = Presentation(str(file_path))
        except Exception as e:
            logger.error(f"Failed to open PPTX file {file_path}: {str(e)}")
            raise HTTPException(
                status_code=400,
                detail="Unable to parse PowerPoint file. The file may be corrupted."
            )

        pages: List[PageContent] = []
        for slide_idx, slide in enumerate(prs.slides):
            slide_texts = []
            
            # Extract text from shapes, text boxes, tables, and images
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        p_text = "".join(run.text for run in paragraph.runs).strip()
                        if p_text:
                            slide_texts.append(p_text)
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            slide_texts.append(row_text)
                
                # Check for picture/image shapes to perform OCR
                if hasattr(shape, "image") and shape.image:
                    try:
                        img_bytes = shape.image.blob
                        img_ocr = run_ocr_on_bytes(img_bytes)
                        cleaned_ocr = self.clean_text(img_ocr)
                        if cleaned_ocr:
                            slide_texts.append(f"[Diagram / Image Text]: {cleaned_ocr}")
                    except Exception as img_err:
                        logger.warning(f"Slide {slide_idx+1} image OCR skipped: {img_err}")

            # Extract slide notes if any
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    slide_texts.append(f"[Notes: {notes}]")

            raw_text = "\n".join(slide_texts)
            cleaned_text = self.clean_text(raw_text)
            
            slide_number = slide_idx + 1
            pages.append(
                PageContent(
                    page_number=slide_number,
                    text=cleaned_text,
                    char_count=len(cleaned_text)
                )
            )

        return pages

    def extract_text(self, file_path: Path, document_id: str, filename: str) -> PDFExtractionResult:
        """
        Extract text from PDF, PowerPoint, or Image files with OCR support.
        """
        if not file_path.exists():
            raise HTTPException(status_code=404, detail="File not found on server.")

        ext = file_path.suffix.lower()
        if ext in [".pptx", ".ppt"]:
            extracted_pages = self._extract_pptx(file_path)
        elif ext in [".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"]:
            extracted_pages = self._extract_image_file(file_path)
        else:
            extracted_pages = self._extract_pdf(file_path)

        total_characters = sum(p.char_count for p in extracted_pages)

        if total_characters == 0 or all(len(p.text.strip()) == 0 for p in extracted_pages):
            raise HTTPException(
                status_code=400,
                detail="The uploaded document contains no extractable text or recognized image text."
            )

        logger.info(
            f"Successfully extracted {len(extracted_pages)} slides/pages ({total_characters} characters) "
            f"from document: {filename} [ID: {document_id}]"
        )

        return PDFExtractionResult(
            document_id=document_id,
            filename=filename,
            total_pages=len(extracted_pages),
            total_characters=total_characters,
            pages=extracted_pages,
            message="Document text & OCR content extracted successfully"
        )

pdf_service = DocumentParserService()
